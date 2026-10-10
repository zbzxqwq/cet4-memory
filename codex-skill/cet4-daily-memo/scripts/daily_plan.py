#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
四级记忆 · 每日单词计划生成器
--------------------------------------------------------
作用：
  1. 读取词库 words.js 和学习进度 进度.json
  2. 按“高频优先 + 错词强化 + 已掌握不再出现”排出今天的 N 个单词
  3. 自动写出两个 Word 表格：
       ../抄写/M.D.docx   —— 只有单词（照抄用）
       ../默写/M.D.docx   —— 单词 + 中文（对答案用）
  4. 在屏幕上打印今天的计划，以及“还没学完”的提醒文字

用法：
  python daily_plan.py                # 按配置生成今天的计划
  python daily_plan.py --count 30     # 今天只要 30 个
  python daily_plan.py --date 10.9    # 指定日期
  python daily_plan.py --dry-run      # 只看计划，不写文件
  python daily_plan.py --status       # 只看进度，不生成
"""

import argparse
import datetime as dt
import json
import os
import random
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
DEFAULT_CONFIG = os.path.join(SKILL_DIR, "config.json")


# --------------------------------------------------------------- 配置
def load_config(path):
    if not os.path.exists(path):
        raise SystemExit(
            "找不到配置文件：%s\n请先按 SKILL.md 的说明创建 config.json。" % path
        )
    with open(path, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    # 相对路径一律相对配置文件所在目录
    base = os.path.dirname(os.path.abspath(path))
    for k in ("dbPath", "progressPath", "copyDir", "dictationDir", "correctionDir"):
        if k in cfg and cfg[k] and not os.path.isabs(cfg[k]):
            cfg[k] = os.path.normpath(os.path.join(base, cfg[k]))
    cfg.setdefault("dailyCount", 50)
    return cfg


# --------------------------------------------------------------- 词库
def parse_words(db_path):
    """解析 words.js 里的 `单词|等级|词性释义` 行。"""
    with open(db_path, "r", encoding="utf-8") as f:
        raw = f.read()
    m = re.search(r"window\.CET4_RAW\s*=\s*`(.*?)`", raw, re.S)
    if not m:
        raise SystemExit("在 %s 里找不到 window.CET4_RAW 词库内容。" % db_path)

    best = {}
    order = []
    for line in m.group(1).split("\n"):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        parts = line.split("|")
        if len(parts) < 3:
            continue
        word = parts[0].strip()
        if not word:
            continue
        try:
            level = int(parts[1])
        except ValueError:
            level = 3
        if level < 1 or level > 4:
            level = 3
        entry = {"word": word, "level": level, "def": parts[2].strip()}
        # 同一个词出现多次时，保留词性/释义更全的那条
        old = best.get(word)
        if old is None:
            best[word] = entry
            order.append(word)
        elif _score(entry["def"]) > _score(old["def"]):
            best[word] = entry
    return [best[w] for w in order]


def _score(text):
    return len(re.findall(r"[a-zA-Z]{1,6}\.", text)) * 1000 + len(text)


def clean_def(text):
    """把 `n.含义1，含义2/v.含义3，含义4` 变成好读的中文串。"""
    out = []
    for chunk in text.split("/"):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"^([a-zA-Z]{1,6}\.)\s*(.*)$", chunk)
        pos, cn = (m.group(1), m.group(2)) if m else ("", chunk)
        out.append((pos + cn).strip())
    return "；".join(out)


# --------------------------------------------------------------- 进度
def load_progress(path):
    if not os.path.exists(path):
        return {"mastered": [], "wrong": {}, "plans": {}, "days": []}
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    # 兼容程序里「导出学习记录」得到的文件（字段名是 masteredWords / wrongWords）
    if not data.get("mastered") and data.get("masteredWords"):
        data["mastered"] = data["masteredWords"]
    if not data.get("wrong") and data.get("wrongWords"):
        data["wrong"] = data["wrongWords"]
    data.setdefault("mastered", [])
    data.setdefault("wrong", {})
    data.setdefault("plans", {})
    data.setdefault("days", [])
    return data


def save_progress(path, data):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# --------------------------------------------------------------- 排计划
def build_plan(words, progress, date_key, count, seed_extra=""):
    mastered = set(progress.get("mastered", []))
    wrong = progress.get("wrong", {})

    due, fresh = [], []
    for e in words:
        if e["word"] in mastered:          # 已掌握：不再出现
            continue
        if wrong.get(e["word"], 0) > 0:    # 不会的词：优先强化
            due.append(e)
        else:
            fresh.append(e)

    due.sort(key=lambda e: -wrong.get(e["word"], 0))

    rnd = random.Random("%s|%s|%d" % (date_key, seed_extra, count))
    groups = {1: [], 2: [], 3: [], 4: []}
    for e in fresh:
        groups[e["level"]].append(e)
    ordered = []
    for lv in (1, 2, 3, 4):                # 考试重点 / 高频词排在最前面
        g = groups[lv]
        rnd.shuffle(g)
        ordered.extend(g)

    plan, seen = [], set()
    for e in due + ordered:
        if len(plan) >= count:
            break
        if e["word"] in seen:
            continue
        seen.add(e["word"])
        plan.append(e)
    return plan


# --------------------------------------------------------------- Word
def write_docx(path, title, subtitle, items, with_meaning):
    try:
        import docx
    except ImportError:
        raise SystemExit("需要 python-docx：python -m pip install python-docx")

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)

    d = docx.Document()
    h = d.add_paragraph()
    h.alignment = 1
    run = h.add_run(title)
    run.bold = True
    run.font.size = docx.shared.Pt(16)

    sub = d.add_paragraph()
    sub.alignment = 1
    sr = sub.add_run(subtitle)
    sr.font.size = docx.shared.Pt(9)
    sr.font.color.rgb = docx.shared.RGBColor(0x66, 0x66, 0x66)

    # 英文单词永远放在最左边一列
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    head = table.rows[0].cells
    head[0].text = "单词"
    head[1].text = "中文意思（参考）" if with_meaning else "抄写一遍"

    for e in items:
        row = table.add_row().cells
        row[0].text = e["word"]
        row[1].text = clean_def(e["def"]) if with_meaning else ""

    d.save(path)
    return path


def month_day_key(when):
    """和已有文件一致：9.5 / 10.8（月.日，不带前导零）"""
    return "%d.%d" % (when.month, when.day)


# --------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser(description="生成四级记忆的每日单词计划")
    ap.add_argument("--config", default=DEFAULT_CONFIG)
    ap.add_argument("--count", type=int, default=None, help="今天学多少个单词")
    ap.add_argument("--date", default=None, help="日期，格式 10.9")
    ap.add_argument("--dry-run", action="store_true", help="只显示，不写文件")
    ap.add_argument("--status", action="store_true", help="只显示学习进度")
    ap.add_argument("--force", action="store_true",
                    help="已经存在同名 Word 文件时也覆盖（默认不覆盖）")
    args = ap.parse_args()

    cfg = load_config(args.config)
    words = parse_words(cfg["dbPath"])
    progress = load_progress(cfg["progressPath"])

    if args.status:
        mastered = progress.get("mastered", [])
        wrong = progress.get("wrong", {})
        left = len([w for w in words if w["word"] not in set(mastered)])
        print("词库总数：%d" % len(words))
        print("已掌握  ：%d" % len(mastered))
        print("还没掌握：%d" % left)
        print("错词数量：%d" % len([w for w, c in wrong.items() if c > 0]))
        if progress.get("days"):
            print("学习天数：%d（最近 %s）" % (len(progress["days"]), progress["days"][-1]))
        return 0

    count = args.count or cfg.get("dailyCount") or 50
    if args.date:
        mm, dd = args.date.split(".")
        today = dt.date(dt.date.today().year, int(mm), int(dd))
    else:
        today = dt.date.today()
    date_key = today.isoformat()

    plan = build_plan(words, progress, date_key, count)
    if not plan:
        print("🎉 词库里的单词已经全部掌握啦，没有新词要学。")
        return 0

    print("=" * 56)
    print("四级记忆 · %s 今日计划（%d 个单词）" % (date_key, len(plan)))
    print("=" * 56)
    for i, e in enumerate(plan, 1):
        flag = "🔁" if progress.get("wrong", {}).get(e["word"]) else "  "
        print("%s %2d. %-18s %s" % (flag, i, e["word"], clean_def(e["def"])))
    print("-" * 56)
    print("🔁 = 之前点过“还不会”的词，今天优先复习。")

    if args.dry_run:
        print("（--dry-run：没有写文件）")
        return 0

    key = month_day_key(today)
    copy_path = os.path.join(cfg["copyDir"], key + ".docx")
    dict_path = os.path.join(cfg["dictationDir"], key + ".docx")

    existing = [p for p in (copy_path, dict_path) if os.path.exists(p)]
    if existing and not args.force:
        print()
        print("⚠️  下面这些文件已经存在，为避免覆盖你写好的内容，本次没有改写：")
        for p in existing:
            print("    " + p)
        print("如果确定要覆盖，请加 --force 再运行一次。")
        return 2

    subtitle = "共 %d 个单词　·　%s　·　四级记忆" % (len(plan), date_key)

    write_docx(copy_path, "四级单词抄写表", subtitle, plan, with_meaning=False)
    write_docx(dict_path, "四级单词默写表", subtitle, plan, with_meaning=True)

    progress.setdefault("plans", {})[date_key] = [e["word"] for e in plan]
    days = progress.setdefault("days", [])
    if date_key not in days:
        days.append(date_key)
    save_progress(cfg["progressPath"], progress)

    print()
    print("✅ 已生成抄写表：%s" % copy_path)
    print("✅ 已生成默写表：%s" % dict_path)
    print("✅ 学习进度已更新：%s" % cfg["progressPath"])
    print()
    print("提醒：今天要完成 %d 个单词。写完记得回到“四级记忆”程序，"
          "把会的点“我会了”，不会的点“还不会”。" % len(plan))
    return 0


if __name__ == "__main__":
    sys.exit(main())
